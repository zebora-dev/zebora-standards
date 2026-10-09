#!/usr/bin/env python3
"""Deck-refresh helpers. Subcommands:

  decode      <drive_tool_result.json> <out_file>     Drive download_file_content result -> real file
  dump        <deck.pptx> [--images DIR]              every text run, table cell, note and image, per slide
  render      <deck.pdf> <prefix> [--dpi 100]         one PNG per slide + 2-up review grids
  textcheck   <original.pptx> <new.pptx>              word-for-word copy check, slide by slide
  livecheck   <live.pptx> <last_build.pptx>           has anyone hand-edited the live deck since the build?
  extract     <deck.pptx> <slide_no> <out.pptx>       one-slide carrier file (for pasting a single slide in)

Needs: python-pptx, PyMuPDF (fitz), Pillow.
"""
import argparse, base64, json, sys
from collections import Counter


def _pres(path):
    from pptx import Presentation
    return Presentation(path)


def decode(a):
    with open(a.src) as f:
        blob = json.load(f)['content']
    with open(a.out, 'wb') as f:
        f.write(base64.b64decode(blob))
    print('wrote', a.out)


def dump(a):
    import os
    p = _pres(a.deck)
    print('canvas', p.slide_width / 914400, 'x', p.slide_height / 914400, 'in')
    if a.images:
        os.makedirs(a.images, exist_ok=True)

    def walk(shapes, i, d=0):
        pad = '  ' * d
        for s in shapes:
            if s.shape_type == 6:
                print(pad + 'GROUP'); walk(s.shapes, i, d + 1); continue
            pos = f'{s.left/914400:.2f},{s.top/914400:.2f} {s.width/914400:.2f}x{s.height/914400:.2f}' if s.left is not None else ''
            if s.shape_type == 13:
                rid = s._element.blipFill.blip.rEmbed
                note = ''
                if a.images:
                    fn = f'{a.images}/s{i:02d}_{rid}.{s.image.ext}'
                    open(fn, 'wb').write(s.image.blob); note = fn
                print(pad + f'PIC {pos} {s.image.size} {note}')
            elif s.shape_type == 19:
                print(pad + f'TABLE {pos}')
                for r in s.table.rows:
                    print(pad + '   |' + ' ¦ '.join(repr(c.text) for c in r.cells))
            elif s.has_text_frame and s.text_frame.text.strip():
                print(pad + f'TXT {pos}')
                for para in s.text_frame.paragraphs:
                    runs = []
                    for r in para.runs:
                        f = ('B' if r.font.bold else '') + ('I' if r.font.italic else '')
                        try:
                            if r.font.color and r.font.color.type == 1: f += str(r.font.color.rgb)
                        except Exception:
                            pass
                        runs.append(f'{r.text!r}<{f}>')
                    if runs: print(pad + '   |' + ' ¦ '.join(runs))
            else:
                print(pad + f'SHP {s.shape_type} {pos}')

    for i, sl in enumerate(p.slides, 1):
        print(f'\n===== SLIDE {i}')
        walk(sl.shapes, i)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            print('NOTES:', repr(sl.notes_slide.notes_text_frame.text))
    print('\nNote: \\x0b in a run is a soft line break; auto-numbered lists show no number in the text.')


def render(a):
    import fitz
    from PIL import Image
    d = fitz.open(a.pdf); ims = []
    for i, pg in enumerate(d):
        f = f'{a.prefix}_s{i+1:02d}.png'; pg.get_pixmap(dpi=a.dpi).save(f); ims.append(Image.open(f))
    w, h = ims[0].size
    for k in range(0, len(ims), 6):
        part = ims[k:k + 6]
        g = Image.new('RGB', (w * 2, h * ((len(part) + 1) // 2)), 'white')
        for i, im in enumerate(part): g.paste(im, ((i % 2) * w, (i // 2) * h))
        g.save(f'{a.prefix}_grid{k // 6 + 1}.png')
    print(len(ims), 'slides ->', f'{a.prefix}_sNN.png and {a.prefix}_gridN.png')


def _words(path):
    out = []
    for sl in _pres(path).slides:
        txt = []
        def walk(shapes):
            for s in shapes:
                if s.shape_type == 6: walk(s.shapes)
                elif s.shape_type == 19:
                    for r in s.table.rows:
                        for c in r.cells: txt.append(c.text)
                elif s.has_text_frame: txt.append(s.text_frame.text)
        walk(sl.shapes)
        out.append(' '.join(' '.join(txt).split()))
    return out


def textcheck(a):
    x, y = _words(a.original), _words(a.new); ok = True
    for i in range(max(len(x), len(y))):
        ca = Counter((x[i] if i < len(x) else '').split()); cb = Counter((y[i] if i < len(y) else '').split())
        miss, extra = ca - cb, cb - ca
        if miss or extra:
            ok = False; print(f'slide {i+1}: missing={dict(miss)} extra={dict(extra)}')
    print('slides', len(x), len(y), 'IDENTICAL' if ok else 'DIFF — every "missing" must be zero; "extra" should only be page numbers')
    # Word multiset per slide: catches dropped/changed/added words, not reordering. Read the slide to check order.


def _shapes(path):
    out = []
    for sl in _pres(path).slides:
        items = []
        def walk(shapes):
            for s in shapes:
                if s.shape_type == 6: walk(s.shapes); continue
                if s.left is None: continue
                t = ' '.join(s.text_frame.text.split()) if s.has_text_frame else ''
                items.append((round(s.left / 914400, 2), round(s.top / 914400, 2), round(s.width / 914400, 2), round(s.height / 914400, 2), t))
        walk(sl.shapes); out.append(items)
    return out


def livecheck(a):
    x, y = _shapes(a.live), _shapes(a.build); bad = 0
    near = lambda p, q: all(abs(p[k] - q[k]) <= 0.03 for k in range(4)) and p[4] == q[4]
    for i in range(max(len(x), len(y))):
        lx = x[i] if i < len(x) else []; ly = y[i] if i < len(y) else []
        ux = [p for p in lx if not any(near(p, q) for q in ly)]; uy = [q for q in ly if not any(near(p, q) for p in lx)]
        if ux or uy:
            bad += 1; print(i + 1, 'LIVE-ONLY', ux, 'BUILD-ONLY', uy)
    print('slides', len(x), len(y), 'edited slides:', bad)
    if bad: print('The live deck has been hand-edited. Do NOT replace the file — paste single slides in, and mirror these edits in the build script.')


def extract(a):
    p = _pres(a.deck); ids = p.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(ids)))):
        if i != a.slide - 1:
            p.part.drop_rel(sid.rId); ids.remove(sid)
    p.save(a.out); print('wrote', a.out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    s = sub.add_parser('decode'); s.add_argument('src'); s.add_argument('out'); s.set_defaults(fn=decode)
    s = sub.add_parser('dump'); s.add_argument('deck'); s.add_argument('--images'); s.set_defaults(fn=dump)
    s = sub.add_parser('render'); s.add_argument('pdf'); s.add_argument('prefix'); s.add_argument('--dpi', type=int, default=100); s.set_defaults(fn=render)
    s = sub.add_parser('textcheck'); s.add_argument('original'); s.add_argument('new'); s.set_defaults(fn=textcheck)
    s = sub.add_parser('livecheck'); s.add_argument('live'); s.add_argument('build'); s.set_defaults(fn=livecheck)
    s = sub.add_parser('extract'); s.add_argument('deck'); s.add_argument('slide', type=int); s.add_argument('out'); s.set_defaults(fn=extract)
    a = ap.parse_args(); a.fn(a)


if __name__ == '__main__':
    main()
